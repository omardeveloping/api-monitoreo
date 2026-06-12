import base64
from concurrent.futures import ThreadPoolExecutor, as_completed
import datetime
import json
import os
import re
import shutil
import socket
import subprocess
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import http.cookiejar
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from django.conf import settings
from openpyxl.styles import Font
from rest_framework.exceptions import ValidationError


TIPOS_ALARMA = {
    "1": "Exceso velocidad",
    "2": "Fatiga conductor",
    "3": "Emergencia SOS",
    "4": "Falla antena GNSS",
    "5": "Cortocircuito GNSS",
    "6": "Voltaje bajo",
    "7": "Apagado",
    "8": "Falla LCD",
    "9": "Alarma TTS",
    "10": "Falla camara",
    "11": "Falla IC",
    "12": "Vel. en zona",
    "13": "Entrada zona",
    "14": "Salida zona",
    "15": "Accidente",
    "16": "Impacto",
    "17": "Freno brusco",
    "18": "Giro brusco",
    "19": "Aceleracion brusca",
    "20": "Colision frontal",
    "21": "Somnolencia",
    "22": "Distraccion",
    "23": "Uso telefono",
    "24": "Fumando",
    "25": "Sin cinturon",
    "26": "Cambio carril",
    "27": "Seguimiento cercano",
    "28": "Zona escolar",
    "29": "Salida ruta",
    "30": "Conduccion nocturna",
}

_NET_MAP = {
    0: "Sin red",
    1: "GPRS",
    2: "3G",
    3: "4G",
    4: "WiFi",
    5: "5G",
    6: "4G+",
    7: "4G+",
    8: "LTE",
    9: "5G",
}
_AES_PARTS = ["A", "B", "c", "D", "e", "F", "g", "H", "I", "J", "k", "L", "m", "n", "O", "P", "Q", "R", "s", "T"]
AES_KEY = ("ttx123456" + _AES_PARTS[0] + _AES_PARTS[4] + _AES_PARTS[18] + "1234").encode()


def _setting(name: str, default=None):
    return getattr(settings, name, os.environ.get(name, default))


def _setting_int(name: str, default: int, *, minimum: int | None = None) -> int:
    try:
        value = int(_setting(name, default))
    except (TypeError, ValueError):
        value = default
    if minimum is not None:
        value = max(minimum, value)
    return value


def _setting_float(name: str, default: float, *, minimum: float | None = None) -> float:
    try:
        value = float(_setting(name, default))
    except (TypeError, ValueError):
        value = default
    if minimum is not None:
        value = max(minimum, value)
    return value


@dataclass(frozen=True)
class CMSV6Config:
    base_url: str
    account: str
    password: str
    device_id: str
    output_dir: str
    canales: int = -1
    task_prepare_wait_secs: int = 180
    task_poll_interval_secs: int = 10
    min_speed_kbps: float = 0.0
    min_speed_window_secs: int = 180
    max_video_download_secs: int = 43200
    downurl_stall_secs: int = 900
    playback_stall_secs: int = 900
    url_rounds: int = 8
    url_round_wait_secs: int = 300
    url_exhaust_wait_secs: int = 1800
    url_exhaust_max_waits: int = 0
    video_search_wait_secs: int = 300
    test_30d_day_retries: int = 0
    test_30d_day_retry_wait_secs: int = 10
    test_30d_scan_newest_first: bool = True
    small_response_bytes: int = 4096
    download_workers: int = 1
    serial_downloads_per_device: bool = True

    @classmethod
    def from_settings(cls, output_dir: str | None = None):
        default_output = (
            output_dir
            or _setting("CMSV6_OUTPUT_DIR", "")
            or _setting("VIDEOS_MDVR_DIR", "")
            or _setting("VIDEOS_IMPORT_DIR", "")
            or str(Path(settings.BASE_DIR) / "cmsv6_output")
        )
        return cls(
            base_url=str(_setting("CMSV6_BASE_URL", "") or "").rstrip("/"),
            account=str(_setting("CMSV6_ACCOUNT", "") or ""),
            password=str(_setting("CMSV6_PASSWORD", "") or ""),
            device_id=str(_setting("CMSV6_DEVICE_ID", "") or ""),
            output_dir=str(default_output),
            canales=_setting_int("CMSV6_CANALES", -1),
            task_prepare_wait_secs=_setting_int("CMSV6_TASK_PREPARE_WAIT_SECS", 180, minimum=0),
            task_poll_interval_secs=_setting_int("CMSV6_TASK_POLL_INTERVAL_SECS", 10, minimum=1),
            min_speed_kbps=_setting_float("CMSV6_MIN_SPEED_KBPS", 0.0, minimum=0.0),
            min_speed_window_secs=_setting_int("CMSV6_MIN_SPEED_WINDOW_SECS", 180, minimum=1),
            max_video_download_secs=_setting_int("CMSV6_MAX_VIDEO_DOWNLOAD_SECS", 43200, minimum=1),
            downurl_stall_secs=_setting_int("CMSV6_DOWNURL_STALL_SECS", 900, minimum=1),
            playback_stall_secs=_setting_int("CMSV6_PLAYBACK_STALL_SECS", 900, minimum=1),
            url_rounds=_setting_int("CMSV6_URL_ROUNDS", 8, minimum=1),
            url_round_wait_secs=_setting_int("CMSV6_URL_ROUND_WAIT_SECS", 300, minimum=0),
            url_exhaust_wait_secs=_setting_int("CMSV6_URL_EXHAUST_WAIT_SECS", 1800, minimum=0),
            url_exhaust_max_waits=_setting_int("CMSV6_URL_EXHAUST_MAX_WAITS", 0, minimum=0),
            video_search_wait_secs=_setting_int("CMSV6_VIDEO_SEARCH_WAIT_SECS", 300, minimum=1),
            test_30d_day_retries=_setting_int("CMSV6_TEST_30D_DAY_RETRIES", 0, minimum=0),
            test_30d_day_retry_wait_secs=_setting_int(
                "CMSV6_TEST_30D_DAY_RETRY_WAIT_SECS", 10, minimum=0
            ),
            test_30d_scan_newest_first=_setting_int("CMSV6_TEST_30D_SCAN_NEWEST_FIRST", 1) != 0,
            small_response_bytes=_setting_int("CMSV6_SMALL_RESPONSE_BYTES", 4096, minimum=1),
            download_workers=_setting_int("CMSV6_DOWNLOAD_WORKERS", 1, minimum=1),
            serial_downloads_per_device=_setting_int(
                "CMSV6_SERIAL_DOWNLOADS_PER_DEVICE",
                1,
                minimum=0,
            )
            != 0,
        )

    def validate(self):
        faltantes = []
        if not self.base_url:
            faltantes.append("CMSV6_BASE_URL")
        if not self.account:
            faltantes.append("CMSV6_ACCOUNT")
        if not self.password:
            faltantes.append("CMSV6_PASSWORD")
        if not self.device_id:
            faltantes.append("CMSV6_DEVICE_ID")
        if not self.output_dir:
            faltantes.append("CMSV6_OUTPUT_DIR")
        if faltantes:
            raise ValidationError(f"Faltan variables CMSV6: {', '.join(faltantes)}.")


class StalledDownloadError(Exception):
    pass


class SlowDownloadError(Exception):
    pass


class DownloadTimeLimitError(Exception):
    pass


class CMSV6AuthError(Exception):
    pass


def _result_code(result) -> int | None:
    if not isinstance(result, dict):
        return None
    try:
        return int(result.get("result"))
    except (TypeError, ValueError):
        return None


def _cmsv6_url_int_param(url: str, name: str) -> int | None:
    try:
        query = urllib.parse.urlparse(url).query
        for key, value in urllib.parse.parse_qsl(query, keep_blank_values=True):
            if key.upper() == name.upper():
                return int(value)
    except (TypeError, ValueError):
        return None
    return None


def _cmsv6_url_reset_foffset(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    if not parsed.query:
        return url
    pairs = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    changed = False
    updated = []
    for key, value in pairs:
        if key.upper() == "FOFFSET":
            value = "0"
            changed = True
        updated.append((key, value))
    if not changed:
        return url
    return urllib.parse.urlunparse(parsed._replace(query=urllib.parse.urlencode(updated)))


def _cmsv6_url_completar_params(url: str, params: dict, *, llenar_vacios=()) -> str:
    if not url:
        return url
    parsed = urllib.parse.urlparse(url)
    if not parsed.query:
        return url
    llenar_vacios_norm = {str(key).upper() for key in llenar_vacios}
    pairs = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    index_by_key = {key.upper(): idx for idx, (key, _value) in enumerate(pairs)}
    changed = False
    for key, value in params.items():
        if value in (None, ""):
            continue
        key_norm = str(key).upper()
        value = str(value)
        if key_norm in index_by_key:
            idx = index_by_key[key_norm]
            current_key, current_value = pairs[idx]
            if current_value == "" and key_norm in llenar_vacios_norm:
                pairs[idx] = (current_key, value)
                changed = True
            continue
        pairs.append((str(key), value))
        changed = True
    if not changed:
        return url
    return urllib.parse.urlunparse(parsed._replace(query=urllib.parse.urlencode(pairs)))


def _crypto_imports():
    try:
        from Crypto.Cipher import AES
        from Crypto.Util.Padding import pad, unpad
    except ImportError as exc:
        raise RuntimeError("Instala pycryptodome para usar CMSV6.") from exc
    return AES, pad, unpad


def _enc(value: str) -> str:
    AES, pad, _unpad = _crypto_imports()
    cipher = AES.new(AES_KEY, AES.MODE_ECB)
    return base64.b64encode(cipher.encrypt(pad(value.encode("utf-8"), AES.block_size))).decode()


def _dec(value: str) -> str:
    AES, _pad, unpad = _crypto_imports()
    data = base64.b64decode(value.replace(" ", "+"))
    return unpad(AES.new(AES_KEY, AES.MODE_ECB).decrypt(data), AES.block_size).decode("utf-8")

def _parse_response(response):
    if isinstance(response, dict) and response.get("encry") == 1 and "data" in response:
        try:
            return json.loads(_dec(response["data"]))
        except Exception:
            pass
    return response


def _ffmpeg_exe() -> str:
    return shutil.which("ffmpeg") or "ffmpeg"


def _ffprobe_exe() -> str:
    return shutil.which("ffprobe") or "ffprobe"


RAW_H264_MIN_FPS = 5.0
RAW_H264_MAX_FPS = 30.0


def is_mp4(path: Path) -> bool:
    try:
        with open(path, "rb") as archivo:
            data = archivo.read(12)
        return len(data) >= 8 and data[4:8] == b"ftyp"
    except OSError:
        return False


def is_h264(path: Path) -> bool:
    try:
        with open(path, "rb") as archivo:
            header = archivo.read(4)
        return header[:4] == b"\x00\x00\x00\x01" or header[:3] == b"\x00\x00\x01"
    except OSError:
        return False

def _mp4_duration_secs(path: Path):
    try:
        result = subprocess.run(
            [
                _ffprobe_exe(),
                "-v",
                "error",
                "-show_entries",
                "format=duration",
                "-of",
                "default=noprint_wrappers=1:nokey=1",
                str(path),
            ],
            capture_output=True,
            timeout=20,
            check=False,
        )
        if result.returncode != 0:
            return None
        value = result.stdout.decode("utf-8", "replace").strip()
        return float(value) if value else None
    except Exception:
        return None


def _count_raw_h264_frames(path: Path) -> int:
    variantes = (
        ["-f", "h264"],
        [],
    )
    for input_args in variantes:
        try:
            result = subprocess.run(
                [
                    _ffprobe_exe(),
                    "-v",
                    "error",
                    *input_args,
                    "-count_frames",
                    "-select_streams",
                    "v:0",
                    "-show_entries",
                    "stream=nb_read_frames,nb_frames",
                    "-of",
                    "json",
                    str(path),
                ],
                capture_output=True,
                timeout=120,
                check=False,
            )
        except (subprocess.TimeoutExpired, FileNotFoundError):
            continue
        if result.returncode != 0:
            continue
        try:
            data = json.loads(result.stdout.decode("utf-8", "replace") or "{}")
        except json.JSONDecodeError:
            continue
        stream = next(iter(data.get("streams") or []), {})
        for key in ("nb_read_frames", "nb_frames"):
            try:
                frames = int(stream.get(key) or 0)
            except (TypeError, ValueError):
                frames = 0
            if frames > 0:
                return frames
    return 0


def _effective_raw_h264_fps(path: Path, expected_secs) -> tuple[float | None, int]:
    try:
        duracion_esperada = float(expected_secs or 0)
    except (TypeError, ValueError):
        return None, 0
    if duracion_esperada <= 0:
        return None, 0

    frames = _count_raw_h264_frames(path)
    if frames <= 0:
        return None, 0

    fps = frames / duracion_esperada
    if fps < RAW_H264_MIN_FPS or fps > RAW_H264_MAX_FPS:
        return None, frames
    return fps, frames


def _mp4_duration_ok(path: Path, expected_secs, log_fn) -> bool:
    if not expected_secs or expected_secs < 60:
        return True
    duration = _mp4_duration_secs(path)
    if duration is None:
        return True
    minimum = max(10.0, float(expected_secs) * 0.50)
    if duration < minimum:
        log_fn(
            f"    Duracion sospechosa: {duration:.1f}s para clip esperado "
            f"de ~{int(expected_secs)}s; probando otro metodo..."
        )
        return False
    return True


def _duration_close_to_expected(path: Path, expected_secs, *, tolerance_ratio=0.01, tolerance_secs=3.0) -> bool:
    if not expected_secs:
        return True
    duration = _mp4_duration_secs(path)
    if duration is None:
        return False
    tolerance = max(float(tolerance_secs), float(expected_secs) * float(tolerance_ratio))
    return abs(duration - float(expected_secs)) <= tolerance


def _retime_mp4_to_expected_duration(path: Path, expected_secs, log_fn) -> bool:
    if not expected_secs:
        return True

    duration = _mp4_duration_secs(path)
    if duration is None or duration <= 0:
        log_fn("    No se pudo medir duracion para ajustar timing CMSV6.")
        return False

    expected = float(expected_secs)
    if _duration_close_to_expected(path, expected):
        return True

    factor = expected / duration
    tmp_path = path.with_name(f"{path.stem}.retime_tmp{path.suffix}")
    if tmp_path.exists():
        tmp_path.unlink()

    log_fn(
        "    Ajustando duracion a timeline CMSV6: "
        f"{duration:.2f}s -> {expected:.2f}s (factor {factor:.6f})."
    )
    command = [
        _ffmpeg_exe(),
        "-y",
        "-hide_banner",
        "-loglevel",
        "error",
        "-i",
        str(path),
        "-vf",
        f"setpts=PTS*{factor:.12f}",
        "-an",
        "-c:v",
        "libx264",
        "-preset",
        "fast",
        "-crf",
        "18",
        "-pix_fmt",
        "yuv420p",
        "-movflags",
        "+faststart",
        str(tmp_path),
    ]
    try:
        result = subprocess.run(command, capture_output=True, timeout=1800, check=False)
    except subprocess.TimeoutExpired:
        log_fn("    Timeout ajustando duracion a timeline CMSV6.")
        return False
    except FileNotFoundError:
        log_fn("    ffmpeg no encontrado.")
        return False

    if result.returncode != 0 or not tmp_path.exists() or tmp_path.stat().st_size <= 4096:
        log_fn("    Fallo ajuste de duracion a timeline CMSV6.")
        if tmp_path.exists():
            tmp_path.unlink()
        return False

    if not _duration_close_to_expected(tmp_path, expected):
        adjusted = _mp4_duration_secs(tmp_path)
        if adjusted is None:
            log_fn("    Duracion ajustada ilegible.")
        else:
            log_fn(
                "    Duracion ajustada aun no coincide con CMSV6: "
                f"{adjusted:.2f}s vs {expected:.2f}s."
            )
        tmp_path.unlink(missing_ok=True)
        return False

    tmp_path.replace(path)
    return True


def _build_raw_h264_reencode_commands(
    ffmpeg: str,
    src_str: str,
    dst_str: str,
    fps_expr: str,
) -> list[list[str]]:
    common = [
        ffmpeg,
        "-y",
        "-hide_banner",
        "-loglevel",
        "error",
        "-probesize",
        "100M",
        "-analyzeduration",
        "100M",
        "-fflags",
        "+genpts+discardcorrupt",
        "-err_detect",
        "ignore_err",
        "-framerate",
        fps_expr,
        "-i",
        src_str,
        "-an",
        "-c:v",
        "libx264",
        "-preset",
        "fast",
        "-crf",
        "18",
        "-pix_fmt",
        "yuv420p",
        "-r",
        fps_expr,
        "-movflags",
        "+faststart",
        dst_str,
    ]
    return [
        [
            ffmpeg,
            "-y",
            "-hide_banner",
            "-loglevel",
            "error",
            "-probesize",
            "100M",
            "-analyzeduration",
            "100M",
            "-fflags",
            "+genpts+discardcorrupt",
            "-err_detect",
            "ignore_err",
            "-f",
            "h264",
            *common[13:],
        ],
        common,
    ]


def convert_to_mp4(src: Path, dst: Path, log_fn, expected_secs=None, *, raw_h264: bool | None = None) -> bool:
    if dst.exists():
        dst.unlink()

    src = Path(src)
    src_str = str(src)
    dst_str = str(dst)
    ffmpeg = _ffmpeg_exe()
    treat_as_raw_h264 = is_h264(src) if raw_h264 is None else raw_h264
    if treat_as_raw_h264:
        fps_value, frames = _effective_raw_h264_fps(src, expected_secs)
        if fps_value is None:
            fps_value = 25.0
            log_fn("    H264 crudo: reencode con FPS fallback 25.")
        else:
            log_fn(
                "    H264 crudo: reencode con FPS efectivo "
                f"{fps_value:.8f} ({frames} frames / {int(expected_secs)}s CMSV6)."
            )
        fps_expr = f"{fps_value:.8f}".rstrip("0").rstrip(".")
        commands = _build_raw_h264_reencode_commands(ffmpeg, src_str, dst_str, fps_expr)
    else:
        commands = [
            [ffmpeg, "-y", "-i", src_str, "-c", "copy", "-movflags", "+faststart", dst_str],
            [
                ffmpeg,
                "-y",
                "-fflags",
                "+genpts",
                "-i",
                src_str,
                "-vf",
                "setpts=N/(25*TB)",
                "-r",
                "25",
                "-c:v",
                "libx264",
                "-preset",
                "fast",
                "-crf",
                "22",
                "-an",
                "-movflags",
                "+faststart",
                dst_str,
            ],
        ]

    for index, command in enumerate(commands, start=1):
        if dst.exists():
            dst.unlink()
        try:
            result = subprocess.run(command, capture_output=True, timeout=1800, check=False)
            if result.returncode == 0 and dst.exists() and dst.stat().st_size > 4096:
                if treat_as_raw_h264 and expected_secs:
                    if _retime_mp4_to_expected_duration(dst, expected_secs, log_fn):
                        return True
                elif _mp4_duration_ok(dst, expected_secs, log_fn):
                    return True
            if index < len(commands):
                log_fn(f"    Intento ffmpeg {index} fallo; probando metodo {index + 1}...")
        except subprocess.TimeoutExpired:
            log_fn(f"    Timeout en intento ffmpeg {index}.")
        except FileNotFoundError:
            log_fn("    ffmpeg no encontrado.")
            return False
    return False


def _clean_excel(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)[:300]
    return value


def _pick(data, *keys, default=""):
    for key in keys:
        if key in data:
            value = data[key]
            if value not in (None, "", "null"):
                return value
    return default


def _parse_gps_time(value):
    if not value:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            continue
    return None


def _normalizar_coord(raw):
    if raw in (None, "", "null"):
        return ""
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return ""
    if value == 0:
        return ""
    if abs(value) > 1000:
        value = value / 1_000_000.0
    return f"{value:.6f}"


def _normalizar_vel(raw):
    if raw in (None, "", "null"):
        return 0
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return 0
    value = value / 10.0
    return round(value, 1)


def _grados_a_direccion(deg):
    try:
        deg = float(deg)
        nombres = ["Norte", "Noreste", "Este", "Sureste", "Sur", "Suroeste", "Oeste", "Noroeste"]
        return f"{nombres[round(deg / 45) % 8]}({int(deg)})"
    except Exception:
        return str(deg)


def _extraer_red_acc(source):
    red = ""
    acc = None
    signal = ""
    if isinstance(source, dict):
        net = source.get("net")
        if net is not None:
            try:
                net_int = int(net)
            except (TypeError, ValueError):
                net_int = None
            red = _NET_MAP.get(net_int, f"Red{net}") if net_int is not None else str(net)
            if net_int == 0:
                signal = "Sin señal"
            elif net_int is not None and net_int >= 3:
                signal = "Excelente"
            elif net_int is not None:
                signal = "Buena"
        ac = source.get("ac")
        if ac == 1:
            acc = True
        elif ac == 0:
            acc = False
        status = source.get("status", source.get("statusInfo", ""))
        if status:
            red2, acc2, signal2 = _extraer_red_acc(str(status))
            red = red2 or red
            acc = acc2 if acc2 is not None else acc
            signal = signal2 or signal
        return red, acc, signal

    text = str(source) if source else ""
    for part in text.replace(";", ",").split(","):
        value = part.strip()
        lower = value.lower()
        if value in ("4G", "3G", "2G", "WiFi", "WIFI", "EDGE", "GPRS", "LTE", "5G", "4G+"):
            red = value
        elif "sin señal" in lower or "no signal" in lower:
            red = "Sin señal"
        elif "excelente" in lower:
            signal = "Excelente"
        elif "buena" in lower and "señal" in lower:
            signal = "Buena"
        elif ("debil" in lower or "débil" in lower or "weak" in lower) and "señal" in lower:
            signal = "Debil"
        elif "acc activado" in lower or "acc on" in lower:
            acc = True
        elif "acc desactivado" in lower or "acc off" in lower:
            acc = False
    return red, acc, signal


def _construir_status(track):
    parts = []
    net = track.get("net")
    if net is not None:
        try:
            net_int = int(net)
        except (TypeError, ValueError):
            net_int = None
        parts.append(_NET_MAP.get(net_int, f"Red{net}") if net_int is not None else str(net))
    ac = track.get("ac")
    if ac == 1:
        parts.append("ACC activado")
    elif ac == 0:
        parts.append("ACC desactivado")
    status = track.get("status", track.get("statusInfo", ""))
    if status and str(status).strip():
        parts.append(str(status).strip())
    return ",".join(parts)


def detectar_eventos_ruta(tracks, device_id):
    eventos = []
    prev_red = None
    prev_signal = None
    prev_acc = None
    stop_start_dt = None
    stop_start_time = ""
    stop_start_loc = ""
    stop_start_lat = ""
    stop_start_lng = ""

    for track in tracks or []:
        gps_time = track.get("gpsTime", track.get("gt", ""))
        server_time = track.get("serverTime", track.get("rt", ""))
        speed = _normalizar_vel(track.get("speed", track.get("sp", track.get("gpsSpeed", 0))))
        lat = track.get("mlat", track.get("latitude", ""))
        lng = track.get("mlng", track.get("longitude", ""))
        if not lat and track.get("lat"):
            lat = str(round(float(track["lat"]) / 1_000_000, 6))
        if not lng and track.get("lng"):
            lng = str(round(float(track["lng"]) / 1_000_000, 6))
        location = track.get("location", track.get("address", track.get("addr", track.get("ls", ""))))
        dt = _parse_gps_time(gps_time)
        red, acc, signal = _extraer_red_acc(track)

        if red and prev_red is not None and red != prev_red:
            eventos.append(
                {
                    "alarmType": "RED",
                    "gpsTime": gps_time,
                    "serverTime": server_time,
                    "alarmName": f"Cambio de red: {prev_red} -> {red}",
                    "speed": speed,
                    "lat": lat,
                    "lng": lng,
                    "location": location,
                    "status": _construir_status(track),
                    "detalle": f"La red cambio de {prev_red} a {red}",
                }
            )
        if red:
            prev_red = red

        if signal and prev_signal is not None and signal != prev_signal:
            eventos.append(
                {
                    "alarmType": "SEÑAL",
                    "gpsTime": gps_time,
                    "serverTime": server_time,
                    "alarmName": f"Señal: {prev_signal} -> {signal}",
                    "speed": speed,
                    "lat": lat,
                    "lng": lng,
                    "location": location,
                    "status": _construir_status(track),
                    "detalle": f"Calidad de señal cambio de {prev_signal} a {signal}",
                }
            )
        if signal:
            prev_signal = signal

        if acc is not None and prev_acc is not None and acc != prev_acc:
            eventos.append(
                {
                    "alarmType": "ACC_ON" if acc else "ACC_OFF",
                    "gpsTime": gps_time,
                    "serverTime": server_time,
                    "alarmName": "MDVR Encendido (ACC ON)" if acc else "MDVR Apagado (ACC OFF)",
                    "speed": speed,
                    "lat": lat,
                    "lng": lng,
                    "location": location,
                    "status": _construir_status(track),
                    "detalle": "Ignicion activada" if acc else "Ignicion desactivada",
                }
            )
        if acc is not None:
            prev_acc = acc

        if speed <= 5:
            if stop_start_dt is None:
                stop_start_dt = dt
                stop_start_time = gps_time
                stop_start_loc = location
                stop_start_lat = lat
                stop_start_lng = lng
        elif stop_start_dt is not None and dt is not None:
            duration = (dt - stop_start_dt).total_seconds()
            if duration >= 60:
                minutes, seconds = divmod(int(duration), 60)
                hours, minutes = divmod(minutes, 60)
                duration_txt = f"{hours}h {minutes}m {seconds}s" if hours else f"{minutes}m {seconds}s"
                eventos.append(
                    {
                        "alarmType": "DETENCION",
                        "gpsTime": stop_start_time,
                        "serverTime": "",
                        "alarmName": f"Detencion {duration_txt}",
                        "speed": 0,
                        "lat": stop_start_lat,
                        "lng": stop_start_lng,
                        "location": stop_start_loc,
                        "status": "",
                        "detalle": f"Detenido {stop_start_time} -> {gps_time} ({duration_txt})",
                    }
                )
            stop_start_dt = None

    if stop_start_dt is not None and tracks:
        last = tracks[-1]
        last_time = last.get("gpsTime", last.get("gt", ""))
        last_dt = _parse_gps_time(last_time)
        if last_dt:
            duration = (last_dt - stop_start_dt).total_seconds()
            if duration >= 60:
                minutes, seconds = divmod(int(duration), 60)
                hours, minutes = divmod(minutes, 60)
                duration_txt = f"{hours}h {minutes}m {seconds}s" if hours else f"{minutes}m {seconds}s"
                eventos.append(
                    {
                        "alarmType": "DETENCION",
                        "gpsTime": stop_start_time,
                        "serverTime": "",
                        "alarmName": f"Detencion {duration_txt}",
                        "speed": 0,
                        "lat": stop_start_lat,
                        "lng": stop_start_lng,
                        "location": stop_start_loc,
                        "status": "",
                        "detalle": f"Detenido {stop_start_time} -> {last_time} ({duration_txt})",
                    }
                )
    return eventos


COLUMNAS_TRACK = [
    "Número de serie",
    "Hora",
    "Recibir Tiempo",
    "Velocidad(km / h)",
    "Velocidad de la grabadora de unidades(km / h)",
    "Driving direction",
    "Valor límite de velocidad actual del vehículo(km / h)",
    "Tipo de carretera actual",
    "Límite actual de velocidad en carretera(km / h)",
    "Millas(km)",
    "Total vehicle mileage(km)",
    "Longitud y Latitud",
    "Ubicación",
    "Estado",
    "Alarma",
    "Información del controlador",
    "Estado de avance y retroceso",
    "Estado de carga y vacío",
    "Estado sobrecargado vacío",
    "Suplemento",
    "Imagenes descargadas",
]
COLUMNAS_ALARMA = [
    "Número de serie",
    "Hora",
    "Recibir Tiempo",
    "Tipo de alarma",
    "Nombre alarma",
    "Velocidad(km / h)",
    "Longitud y Latitud",
    "Ubicación",
    "Estado",
    "Información adicional",
]
COLUMNAS_EVENTOS = [
    "Número de serie",
    "Hora GPS",
    "Tipo de evento",
    "Descripcion",
    "Velocidad(km / h)",
    "Longitud y Latitud",
    "Ubicacion",
    "Detalle",
]
NOMBRE_EVENTO = {
    "RED": "Cambio de red",
    "SEÑAL": "Cambio de señal",
    "ACC_ON": "MDVR Encendido",
    "ACC_OFF": "MDVR Apagado",
    "DETENCION": "Detencion",
}


def _track_a_fila(index, track, device_id):
    lat = _normalizar_coord(_pick(track, "mlat", "latitude", "lat", "mapLat", "blat"))
    lng = _normalizar_coord(_pick(track, "mlng", "longitude", "lng", "mapLng", "blng"))
    direction = _pick(track, "direction", "hx", "heading", "hd")
    supplement = track.get("supplement", track.get("suplemento", track.get("ifSupplement", 0)))
    return [
        f"{index} {device_id}({device_id})",
        track.get("gpsTime", track.get("gt", track.get("gps_time", ""))),
        track.get("serverTime", track.get("rt", track.get("receive_time", ""))),
        _normalizar_vel(_pick(track, "speed", "sp", "gpsSpeed", "gs")),
        _normalizar_vel(_pick(track, "driverSpeed", "dspeed", "recorderSpeed", "rs", "ds")),
        _grados_a_direccion(direction) if direction != "" else "",
        track.get("speedLimit", track.get("vehicleSpeedLimit", None)),
        track.get("roadType", None),
        track.get("roadSpeedLimit", None),
        _pick(track, "mileage", "mile", "pk", default=0) or 0,
        _pick(track, "totalMileage", "totalMile", "tm", "pk", default=0) or 0,
        f"{lat},{lng}" if lat and lng else "",
        track.get("location", track.get("address", track.get("addr", track.get("ls", "")))),
        _construir_status(track),
        track.get("alarm", track.get("alarmInfo", track.get("alarmType", None))),
        track.get("driverInfo", None),
        track.get("forwardReverseStatus", None),
        track.get("loadEmptyStatus", None),
        track.get("overloadStatus", None),
        "Si" if supplement and str(supplement) not in ("0", "", "None", "False") else "No",
        track.get("downloadedImages", None),
    ]


def _escribir_hoja(wb, title, columns, rows):
    font_std = Font(name="Calibri", size=12)
    ws = wb.create_sheet(title=title)
    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(1, col_idx, str(column))
        cell.font = font_std
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = max(16, len(str(column)) + 3)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row_idx, col_idx, _clean_excel(value))
            cell.font = font_std
    ws.freeze_panes = "A2"
    return ws


def guardar_tracks_json(
    carpeta_base: Path,
    device_id: str,
    fecha: datetime.date,
    tracks: list,
    log_fn,
) -> Path | None:
    if not tracks:
        return None

    carpeta_base.mkdir(parents=True, exist_ok=True)
    path = carpeta_base / f"{device_id} {fecha.isoformat()}_tracks.json"
    payload = {
        "device_id": device_id,
        "fecha": fecha.isoformat(),
        "generado_en": datetime.datetime.now().isoformat(timespec="seconds"),
        "tracks": tracks,
    }
    with open(path, "w", encoding="utf-8") as archivo:
        json.dump(payload, archivo, ensure_ascii=False, default=str)
    log_fn(f"  [OK] Tracks CMSV6 guardados: {path.name} ({len(tracks)} puntos)")
    return path


def exportar_excel(
    carpeta_base: Path,
    device_id: str,
    fecha_ini: datetime.datetime,
    fecha_fin: datetime.datetime,
    gps_actual: dict,
    tracks: list,
    alarmas: list,
    eventos: list,
    log_fn,
    *,
    hacer_ruta=True,
    hacer_alarmas=True,
):
    carpeta_base.mkdir(parents=True, exist_ok=True)
    filename_base = (
        f"{device_id} {fecha_ini.strftime('%Y-%m-%d %H-%M-%S')}"
        f"~{fecha_fin.strftime('%Y-%m-%d %H-%M-%S')}"
    )
    route_path = None

    if hacer_ruta:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        rows = [_track_a_fila(index + 1, track, device_id) for index, track in enumerate(tracks or [])]
        _escribir_hoja(wb, "Track point", COLUMNAS_TRACK, rows)
        route_path = carpeta_base / f"{filename_base}.xlsx"
        wb.save(route_path)
        log_fn(f"  [OK] Ruta GPS guardada: {route_path.name} ({len(rows)} puntos)")

    if hacer_alarmas:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        alarm_rows = []
        for index, alarm in enumerate(alarmas or [], 1):
            alarm_type = str(_pick(alarm, "alarmType", "type", "atype", "alarm", "at", default="?"))
            lat = _normalizar_coord(_pick(alarm, "mlat", "lat", "latitude", "mapLat"))
            lng = _normalizar_coord(_pick(alarm, "mlng", "lng", "longitude", "mapLng"))
            alarm_rows.append(
                [
                    f"{index} {device_id}({device_id})",
                    _pick(alarm, "gpsTime", "gt", "alarmTime", "alarm_time", "bt"),
                    _pick(alarm, "serverTime", "rt", "receive_time", "et"),
                    alarm_type,
                    TIPOS_ALARMA.get(alarm_type, f"Tipo {alarm_type}"),
                    _normalizar_vel(_pick(alarm, "speed", "sp", "gpsSpeed")),
                    f"{lat},{lng}" if lat and lng else "",
                    _pick(alarm, "location", "address", "addr", "ls"),
                    _pick(alarm, "status", "stl", "state"),
                    _clean_excel(alarm),
                ]
            )
        _escribir_hoja(wb, "Alarmas", COLUMNAS_ALARMA, alarm_rows)

        event_rows = []
        for index, event in enumerate(eventos or [], 1):
            lat = event.get("lat", "")
            lng = event.get("lng", "")
            event_rows.append(
                [
                    f"{index} {device_id}({device_id})",
                    event.get("gpsTime", ""),
                    NOMBRE_EVENTO.get(event.get("alarmType", ""), event.get("alarmType", "")),
                    event.get("alarmName", ""),
                    event.get("speed", 0),
                    f"{lat},{lng}" if lat and lng else "",
                    event.get("location", ""),
                    event.get("detalle", ""),
                ]
            )
        _escribir_hoja(wb, "Eventos", COLUMNAS_EVENTOS, event_rows)

        summary_alarm = {}
        for alarm in alarmas or []:
            alarm_type = str(alarm.get("alarmType", alarm.get("type", alarm.get("atype", "?"))))
            summary_alarm[alarm_type] = summary_alarm.get(alarm_type, 0) + 1
        summary_alarm_rows = [
            (alarm_type, TIPOS_ALARMA.get(alarm_type, f"Tipo {alarm_type}"), count)
            for alarm_type, count in sorted(summary_alarm.items(), key=lambda item: -item[1])
        ] or [("-", "Sin alarmas registradas", 0)]
        _escribir_hoja(wb, "Resumen Alarmas", ["Tipo", "Nombre", "Cantidad"], summary_alarm_rows)

        summary_events = {}
        for event in eventos or []:
            event_type = NOMBRE_EVENTO.get(event.get("alarmType", ""), event.get("alarmType", ""))
            summary_events[event_type] = summary_events.get(event_type, 0) + 1
        summary_event_rows = [
            (event_type, count) for event_type, count in sorted(summary_events.items(), key=lambda item: -item[1])
        ] or [("Sin eventos detectados", 0)]
        _escribir_hoja(wb, "Resumen Eventos", ["Tipo de evento", "Cantidad"], summary_event_rows)

        alarm_path = carpeta_base / f"{device_id} {fecha_ini.strftime('%Y-%m-%d')}_Alarmas.xlsx"
        wb.save(alarm_path)
        log_fn(
            f"  [OK] Alarmas+Eventos guardados: {alarm_path.name} "
            f"({len(alarmas or [])} alarmas, {len(eventos or [])} eventos)"
        )
    return route_path


class CMSV6Session:
    def __init__(self, config: CMSV6Config):
        self.config = config
        self._reset_opener()
        self.jsession = None
        self.sid = None
        self.api_login_response = {}
        self.web_login_response = {}

    def _reset_opener(self):
        self.cookie_jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cookie_jar))

    def login(self):
        data = urllib.parse.urlencode(
            {
                "account": self.config.account,
                "password": self.config.password,
                "lang": "0",
                "clientType": "web",
            }
        ).encode()
        request = urllib.request.Request(
            f"{self.config.base_url}/808gps/StandardApiAction_login.action",
            data=data,
            method="POST",
        )
        with self.opener.open(request, timeout=30) as response:
            api_response = json.loads(response.read())
        self.api_login_response = api_response
        self.jsession = api_response.get("jsession")
        if not self.jsession:
            raise Exception(f"Login API sin jsession: {api_response}")

        self.opener.open(f"{self.config.base_url}/808gps/login.html", timeout=30)
        request = urllib.request.Request(
            f"{self.config.base_url}/808gps/StandardLoginAction_initLoginSession.action?{_enc('{}')}",
            data=b"",
            method="POST",
            headers={"Newv": "1"},
        )
        with self.opener.open(request, timeout=30) as response:
            self.sid = json.loads(_dec(json.loads(response.read())["data"]))["jsessionId"]

        password = _enc(base64.b64encode(urllib.parse.quote(self.config.password, safe="").encode()).decode())
        login_payload = {
            "account": self.config.account,
            "ipson": password,
            "language": "es",
            "verificationCode": "",
            "v9OldStyle": "",
        }
        request = urllib.request.Request(
            f"{self.config.base_url}/808gps/StandardLoginAction_login.action?{_enc(json.dumps(login_payload))}",
            data=b"",
            method="POST",
            headers={"Newv": "1", "jsessionId": self.sid},
        )
        with self.opener.open(request, timeout=30) as response:
            result = _parse_response(json.loads(response.read()))
        self.web_login_response = result
        if result.get("result") != 0:
            raise Exception(f"Login CMSV6 fallido: {result}")

    def _post_web(self, endpoint, url_params, body_params, timeout=30):
        url = f"{self.config.base_url}/808gps/{endpoint}?{_enc(json.dumps(url_params))}"
        request = urllib.request.Request(
            url,
            data=_enc(json.dumps(body_params)).encode(),
            method="POST",
            headers={
                "Newv": "1",
                "jsessionId": self.sid,
                "csrfToken": "",
                "Content-Type": "text/plain;charset=UTF-8",
            },
        )
        with self.opener.open(request, timeout=timeout) as response:
            return _parse_response(json.loads(response.read().decode("utf-8")))

    def _post_api(self, endpoint, params, timeout=30):
        params = dict(params)
        params["jsession"] = self.jsession
        data = urllib.parse.urlencode(params).encode()
        request = urllib.request.Request(f"{self.config.base_url}/808gps/{endpoint}", data=data, method="POST")
        with self.opener.open(request, timeout=timeout) as response:
            return _parse_response(json.loads(response.read()))

    def get_gps(self):
        result = self._post_web(
            "StandardPositionAction_statusEx.action",
            {"toMap": "1", "newv": "1"},
            {"devIdnos": self.config.device_id},
        )
        return result.get("status", [])

    def get_track(self, fecha):
        begin = fecha.strftime("%Y-%m-%d") + " 00:00:00"
        end = fecha.strftime("%Y-%m-%d") + " 23:59:59"
        page_size = 5000
        attempts = [
            (
                "StandardApiAction_queryTrackDetail.action",
                {
                    "devIdno": self.config.device_id,
                    "begintime": begin,
                    "endtime": end,
                    "pageRecords": page_size,
                },
            ),
            (
                "StandardApiAction_queryTrackInfo.action",
                {
                    "devIdno": self.config.device_id,
                    "begintime": begin,
                    "endtime": end,
                    "pageRecords": page_size,
                },
            ),
        ]
        for endpoint, params in attempts:
            try:
                rows_total = []
                page = 1
                while page <= 100:
                    page_params = dict(params)
                    page_params["currentPage"] = page
                    result = self._post_api(endpoint, page_params)
                    rows = (
                        result.get("trackDetails")
                        or result.get("tracks")
                        or result.get("rows")
                        or []
                    )
                    if isinstance(rows, list) and rows:
                        rows_total.extend(rows)
                    pagination = result.get("pagination") or {}
                    try:
                        total_pages = int(pagination.get("totalPages") or 0)
                    except (TypeError, ValueError):
                        total_pages = 0
                    has_next = pagination.get("hasNextPage") in (True, "true", "True", 1, "1")
                    if total_pages:
                        has_next = page < total_pages
                    if not has_next or not rows:
                        break
                    page += 1
                if rows_total:
                    return rows_total
            except Exception:
                pass
        return []

    def get_alarms(self, fecha):
        begin = fecha.strftime("%Y-%m-%d") + " 00:00:00"
        end = fecha.strftime("%Y-%m-%d") + " 23:59:59"
        attempts = [
            (
                "web",
                "StandardAlarmAction_queryAlarmInfo.action",
                {"newv": "1"},
                {
                    "devIdno": self.config.device_id,
                    "begintime": begin,
                    "endtime": end,
                    "alarmType": -1,
                    "currentPage": 1,
                    "pageRecords": 5000,
                },
            ),
            (
                "api",
                "StandardApiAction_queryAlarmInfo.action",
                {
                    "devIdno": self.config.device_id,
                    "begintime": begin,
                    "endtime": end,
                    "alarmType": -1,
                    "currentPage": 1,
                    "pageRecords": 5000,
                },
            ),
        ]
        for kind, endpoint, *args in attempts:
            try:
                result = self._post_web(endpoint, args[0], args[1]) if kind == "web" else self._post_api(endpoint, args[0])
                rows = result.get("alarmDetails") or result.get("alarms") or result.get("rows") or result.get("list") or []
                if isinstance(rows, list) and rows:
                    return rows
            except Exception:
                pass
        return []

    def relogin_api(self):
        data = urllib.parse.urlencode(
            {
                "account": self.config.account,
                "password": self.config.password,
                "lang": "0",
                "clientType": "web",
            }
        ).encode()
        request = urllib.request.Request(
            f"{self.config.base_url}/808gps/StandardApiAction_login.action",
            data=data,
            method="POST",
        )
        with self.opener.open(request, timeout=30) as response:
            result = json.loads(response.read())
        self.jsession = result.get("jsession")
        if not self.jsession:
            raise Exception(f"Re-login API sin jsession: {result}")

    def full_relogin(self):
        self._reset_opener()
        self.jsession = None
        self.sid = None
        self.login()

    def _archivo_video_datetime(self, archivo, fecha, key):
        fecha_base = fecha.date() if isinstance(fecha, datetime.datetime) else fecha
        segundos = _segundos_video(archivo, key)
        return datetime.datetime.combine(fecha_base, datetime.time.min) + datetime.timedelta(
            seconds=segundos
        )

    def _download_task_status(self, task):
        for key in ("stu", "status", "st"):
            try:
                return int(task.get(key))
            except (TypeError, ValueError):
                continue
        return None

    def _download_task_matches(self, task, archivo):
        archivo_chn = _video_channel_idx(archivo)
        try:
            task_chn = int(task.get("chn", task.get("channel", task.get("ch", -999))))
        except (TypeError, ValueError):
            task_chn = -999
        if task_chn not in {-999, archivo_chn, archivo_chn + 1}:
            return False

        archivo_nombre = Path(str(archivo.get("file", "") or "")).name
        task_path = str(
            task.get("fph")
            or task.get("filePath")
            or task.get("filename")
            or task.get("file")
            or ""
        )
        task_nombre = Path(task_path).name
        if archivo_nombre and task_nombre and archivo_nombre != task_nombre:
            return False
        return True

    def get_download_tasks(self, archivo, fecha, timeout=30):
        begin = self._archivo_video_datetime(archivo, fecha, "beg")
        end = self._archivo_video_datetime(archivo, fecha, "end")
        if end <= begin:
            end = begin + datetime.timedelta(seconds=1)
        result = self._post_api(
            "StandardApiAction_downloadTasklist.action",
            {
                "devIdno": self.config.device_id,
                "begintime": begin.strftime("%Y-%m-%d %H:%M:%S"),
                "endtime": end.strftime("%Y-%m-%d %H:%M:%S"),
                "currentPage": 1,
                "pageRecords": 50,
            },
            timeout=timeout,
        )
        rows = (
            result.get("infos")
            or result.get("tasks")
            or result.get("rows")
            or result.get("list")
            or []
        )
        return rows if isinstance(rows, list) else []

    def wait_download_task(self, archivo, fecha, log_fn=None):
        log = log_fn or (lambda _msg: None)
        wait_secs = max(0, int(self.config.task_prepare_wait_secs or 0))
        if wait_secs <= 0:
            return None

        deadline = time.monotonic() + wait_secs
        last_status = None
        last_match = None
        log(f"    Esperando tarea CMSV6 ({wait_secs}s max)...")
        while True:
            try:
                tasks = self.get_download_tasks(archivo, fecha, timeout=30)
                matches = [task for task in tasks if self._download_task_matches(task, archivo)]
                if matches:
                    last_match = matches[0]
                    last_status = self._download_task_status(last_match)
                    log(f"    Tarea CMSV6 status={last_status} ({len(matches)} coincidencia/s)")
                    if last_status == 4:
                        return last_match
                    if last_status in {3, 5, 6, 7}:
                        return last_match
            except Exception as exc:
                log(f"    No se pudo consultar downloadTasklist: {exc}")

            if time.monotonic() >= deadline:
                if last_status is not None:
                    log(f"    Tarea CMSV6 sin completar; ultimo status={last_status}")
                return last_match
            time.sleep(min(self.config.task_poll_interval_secs, max(1, int(deadline - time.monotonic()))))

    def get_video_files(self, fecha, log_fn=None):
        log = log_fn or (lambda _msg: None)
        user_agent = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CMSV6Client/4.0"
        fecha_base = fecha.date() if isinstance(fecha, datetime.datetime) else fecha

        def _query(endpoint, *, fileattr=2, rectype=-1, timeout=20, session_kind="api"):
            session_token = self.sid if session_kind == "web" else self.jsession
            params = {
                "DevIDNO": self.config.device_id,
                "LOC": 1,
                "CHN": self.config.canales,
                "YEAR": fecha_base.year,
                "MON": fecha_base.month,
                "DAY": fecha_base.day,
                "RECTYPE": rectype,
                "FILEATTR": fileattr,
                "BEG": 0,
                "END": 86399,
                "jsession": session_token,
                "ARM1": 0,
                "ARM2": 0,
                "RES": 0,
                "STREAM": -1,
                "STORE": 0,
            }
            data = urllib.parse.urlencode(params).encode()
            request = urllib.request.Request(
                f"{self.config.base_url}/808gps/{endpoint}", data=data, method="POST"
            )
            request.add_header("User-Agent", user_agent)
            if session_kind == "web" and self.sid:
                request.add_header("Newv", "1")
                request.add_header("jsessionId", self.sid)
            with self.opener.open(request, timeout=timeout) as response:
                return _parse_response(json.loads(response.read()))

        def _extract_files(result):
            for key in ("files", "fileList", "list", "data", "rows", "result_data"):
                value = result.get(key)
                if isinstance(value, list) and value:
                    return value
                if isinstance(value, dict):
                    for subkey in ("files", "fileList", "list"):
                        subvalue = value.get(subkey)
                        if isinstance(subvalue, list) and subvalue:
                            return subvalue
            return []

        auth_results = (5,)

        def _query_valid(endpoint, *, fileattr=2, rectype=-1, timeout=20):
            session_kinds = ["api"]
            if self.sid:
                session_kinds.append("web")

            last_auth = ""
            for session_kind in session_kinds:
                result = _query(
                    endpoint,
                    fileattr=fileattr,
                    rectype=rectype,
                    timeout=timeout,
                    session_kind=session_kind,
                )
                code = _result_code(result)
                if code not in auth_results:
                    if session_kind == "web":
                        log("  Consulta de video aceptada usando sesion web.")
                    return result
                last_auth = f"{session_kind}: result={code}"

            log(f"  Sesion expirada/rechazada ({last_auth}), renovando...")
            for attempt in range(1, 4):
                try:
                    self.relogin_api() if attempt == 1 else self.full_relogin()
                    session_kinds = ["api"]
                    if self.sid:
                        session_kinds.append("web")
                    for session_kind in session_kinds:
                        result = _query(
                            endpoint,
                            fileattr=fileattr,
                            rectype=rectype,
                            timeout=timeout,
                            session_kind=session_kind,
                        )
                        code = _result_code(result)
                        if code not in auth_results:
                            log(
                                "  Sesion renovada correctamente"
                                + (" usando sesion web." if session_kind == "web" else ".")
                            )
                            return result
                        last_auth = f"{session_kind}: result={code}"
                except Exception as exc:
                    last_auth = str(exc)
                    log(f"  Re-login intento {attempt}/3 fallo: {exc}")
                time.sleep(1)
            raise CMSV6AuthError(f"sesion CMSV6 rechazada ({last_auth})")

        ep1 = "StandardApiAction_getVideoFileInfo.action"
        combos = [(ep1, 2, -1)]
        auth_errors = []
        for endpoint, fileattr, rectype in combos:
            combo = f"LOC=1 FA={fileattr} RT={rectype}"
            try:
                log(f"  Probando {combo}...")
                result = _query_valid(endpoint, fileattr=fileattr, rectype=rectype, timeout=25)
                code = _result_code(result)
                if code == 0:
                    files = _extract_files(result)
                    if files:
                        log(f"  OK: {len(files)} archivos ({combo})")
                        return files
                    log(f"  Sin grabaciones ({combo})")
                    return []
                if code == 32:
                    log(f"  Equipo offline para consulta de video ({combo}); no hay descarga directa MDVR.")
                    return []
                if code == 3:
                    log(f"  Parametros rechazados por CMSV6 ({combo}); revisar LOC/FILEATTR/RECTYPE.")
                    return []
                else:
                    log(f"  result={code} ({combo})")
            except urllib.error.HTTPError as exc:
                log(f"  Error HTTP {exc.code} ({combo})")
            except CMSV6AuthError as exc:
                auth_errors.append(f"{combo}: {exc}")
                log(f"  Error auth ({combo}): {exc}; probando siguiente combinacion...")
            except Exception as exc:
                log(f"  Error consultando {combo}: {exc}")
        if auth_errors:
            resumen = "; ".join(auth_errors[:4])
            extra = "" if len(auth_errors) <= 4 else f"; +{len(auth_errors) - 4} mas"
            raise CMSV6AuthError(f"sesion CMSV6 rechazada en combinaciones de video ({resumen}{extra})")
        return []

    def refresh_url(self, url):
        if not url or not self.jsession:
            return url
        if url.startswith("/"):
            parsed_base = urllib.parse.urlparse(self.config.base_url)
            url = f"{parsed_base.scheme}://{parsed_base.netloc}{url}"
        elif not url.startswith("http"):
            url = f"{self.config.base_url}/{url.lstrip('/')}"
        if "jsession=" in url:
            return re.sub(r"(jsession=)[^&]+", rf"\g<1>{self.jsession}", url)
        separator = "&" if "?" in url else "?"
        return f"{url}{separator}jsession={self.jsession}"

    def _partial_matches_server(self, current_url, dest_path: Path, offset: int, timeout: int, log) -> bool:
        probe_len = min(64 * 1024, offset)
        if probe_len <= 0:
            return True
        probe_start = offset - probe_len
        request = urllib.request.Request(current_url)
        request.add_header("User-Agent", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CMSV6Client/4.0")
        request.add_header("Range", f"bytes={probe_start}-")
        with self.opener.open(request, timeout=timeout) as response:
            status_code = getattr(response, "status", response.getcode())
            if probe_start > 0 and status_code != 206:
                raise Exception("Servidor no acepto validacion HTTP Range del parcial")
            remote_probe = response.read(probe_len)
        if len(remote_probe) != probe_len:
            raise Exception("Servidor devolvio probe incompleto para validar parcial")
        with open(dest_path, "rb") as local_file:
            local_file.seek(probe_start)
            local_probe = local_file.read(probe_len)
        if local_probe == remote_probe:
            log(f"    Parcial validado con HTTP Range ({offset / 1048576:.1f} MB)")
            return True
        log("    Parcial no coincide con servidor; reiniciando ese archivo para evitar duplicados")
        return False

    def download_file(
        self,
        url,
        dest_path,
        *,
        progress_cb=None,
        log_fn=None,
        max_retries=10,
        dl_timeout=120,
        stall_secs=45,
        min_speed_kbps=None,
        min_speed_secs=None,
        max_total_secs=None,
    ):
        log = log_fn or (lambda _msg: None)
        dest_path = Path(dest_path)
        last_exc = None
        min_speed_kbps = self.config.min_speed_kbps if min_speed_kbps is None else float(min_speed_kbps)
        min_speed_secs = self.config.min_speed_window_secs if min_speed_secs is None else int(min_speed_secs)
        max_total_secs = self.config.max_video_download_secs if max_total_secs is None else int(max_total_secs)
        call_start = time.monotonic()

        for attempt in range(1, max_retries + 1):
            try:
                current_url = _cmsv6_url_reset_foffset(self.refresh_url(url))
                expected_total = _cmsv6_url_int_param(current_url, "FLENGTH")
                offset = 0
                if dest_path.exists() and dest_path.stat().st_size >= self.config.small_response_bytes:
                    offset = dest_path.stat().st_size
                    log(f"    Reanudando desde {offset / 1048576:.1f} MB")

                if expected_total and offset > expected_total:
                    log(
                        f"    Parcial excede tamano esperado "
                        f"({offset / 1048576:.1f}/{expected_total / 1048576:.1f} MB); reiniciando"
                    )
                    dest_path.unlink()
                    offset = 0

                if offset > 0:
                    if self._partial_matches_server(current_url, dest_path, offset, min(stall_secs, dl_timeout), log):
                        if expected_total and offset == expected_total:
                            log("    Archivo temporal ya esta completo; no se re-descarga")
                            return offset
                    else:
                        dest_path.unlink()
                        offset = 0

                request = urllib.request.Request(current_url)
                request.add_header("User-Agent", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CMSV6Client/4.0")
                if offset > 0:
                    request.add_header("Range", f"bytes={offset}-")

                socket_timeout = min(stall_secs, dl_timeout)
                with self.opener.open(request, timeout=socket_timeout) as response:
                    status_code = getattr(response, "status", 200)
                    if status_code in (401, 403):
                        raise Exception(f"HTTP {status_code}: sesion expirada")
                    resumed = offset > 0
                    if resumed and status_code != 206:
                        raise Exception("Servidor no acepto reanudacion HTTP Range; parcial conservado")

                    content_length = int(response.headers.get("Content-Length", 0) or 0)
                    total = expected_total or content_length
                    if resumed and content_length > 0 and not expected_total:
                        total = content_length + offset
                    downloaded = offset
                    mode = "ab" if resumed else "wb"
                    last_byte_t = time.monotonic()
                    speed_t = last_byte_t
                    speed_bytes = downloaded
                    last_log_t = last_byte_t
                    with open(dest_path, mode) as output:
                        while True:
                            now = time.monotonic()
                            if max_total_secs and now - call_start > max_total_secs:
                                raise DownloadTimeLimitError(
                                    f"Tiempo maximo {max_total_secs // 60} min superado "
                                    f"({downloaded / 1048576:.1f} MB descargados)"
                                )
                            try:
                                chunk = response.read(65536)
                            except socket.timeout:
                                if time.monotonic() - last_byte_t >= stall_secs:
                                    raise StalledDownloadError(f"Sin datos por {stall_secs}s")
                                continue
                            if not chunk:
                                break
                            if expected_total and downloaded >= expected_total:
                                break
                            reached_expected_total = False
                            if expected_total and downloaded + len(chunk) > expected_total:
                                allowed = expected_total - downloaded
                                if allowed <= 0:
                                    break
                                chunk = chunk[:allowed]
                                reached_expected_total = True
                            output.write(chunk)
                            downloaded += len(chunk)
                            last_byte_t = time.monotonic()
                            elapsed_speed = last_byte_t - speed_t
                            if elapsed_speed >= min_speed_secs:
                                delta = downloaded - speed_bytes
                                kbps = (delta / 1024) / elapsed_speed if elapsed_speed > 0 else 0
                                if min_speed_kbps > 0 and kbps < min_speed_kbps:
                                    raise SlowDownloadError(
                                        f"Velocidad baja {kbps:.1f} KB/s por {int(elapsed_speed)}s"
                                    )
                                speed_t = last_byte_t
                                speed_bytes = downloaded
                            if last_byte_t - last_log_t >= 30:
                                log(f"    {downloaded / 1048576:.1f} MB descargados")
                                last_log_t = last_byte_t
                            if progress_cb:
                                progress_cb(downloaded, total)
                            if reached_expected_total:
                                break
                if downloaded <= offset:
                    raise Exception("El servidor devolvio 0 bytes")
                if expected_total and downloaded > expected_total:
                    raise Exception("Descarga excedio el tamano esperado")
                return downloaded
            except (StalledDownloadError, SlowDownloadError, DownloadTimeLimitError):
                raise
            except urllib.error.HTTPError as exc:
                last_exc = exc
                if exc.code in (400, 404, 405):
                    log(f"    HTTP {exc.code}: URL no valida")
                    break
                if exc.code in (401, 403):
                    try:
                        self.full_relogin()
                    except Exception:
                        pass
                if attempt < max_retries:
                    time.sleep(min(3 * attempt, 15))
            except Exception as exc:
                last_exc = exc
                if attempt < max_retries:
                    log(f"    Reintento {attempt}/{max_retries}: {exc}")
                    try:
                        self.full_relogin()
                    except Exception:
                        pass
                    time.sleep(min(3 * attempt, 15))
        raise last_exc or Exception("Descarga fallida")


_CMSV6_MDVR_LEGACY_NAME_RE = re.compile(
    r"^\d+-(?P<camara>\d{2})-\d{6}-\d{6}-.+\.(?:h264|mp4)$",
    re.IGNORECASE,
)
_CMSV6_MDVR_NEW_NAME_RE = re.compile(
    r"^\d+-\d{6}-\d{6}-\d{6}-(?P<codigo>\d+)\.(?:mp4)$",
    re.IGNORECASE,
)


def _fecha_nombre_video(fecha: datetime.date | datetime.datetime | None) -> datetime.date:
    if isinstance(fecha, datetime.datetime):
        return fecha.date()
    if isinstance(fecha, datetime.date):
        return fecha
    return datetime.datetime.now().date()


def _codigo_video_compatible_mdvr(channel: int) -> str:
    camara = channel + 1 if 0 <= channel <= 3 else channel
    if camara not in {1, 2, 3, 4}:
        camara = 1
    return f"2001{camara:02d}00"


def _cmsv6_segundos_a_hhmmss(valor) -> int:
    try:
        segundos = int(valor or 0)
    except (TypeError, ValueError):
        return 0
    segundos = max(0, segundos) % 86400
    hh = segundos // 3600
    mm = (segundos % 3600) // 60
    ss = segundos % 60
    return int(f"{hh:02d}{mm:02d}{ss:02d}")


def _nombre_compatible_mdvr(nombre: str) -> bool:
    match_legacy = _CMSV6_MDVR_LEGACY_NAME_RE.match(nombre or "")
    if match_legacy:
        try:
            return int(match_legacy.group("camara")) in {1, 2, 3, 4}
        except (TypeError, ValueError):
            return False

    match_nuevo = _CMSV6_MDVR_NEW_NAME_RE.match(nombre or "")
    if not match_nuevo:
        return False
    codigo = match_nuevo.group("codigo")
    if not codigo or len(codigo) < 6:
        return False
    try:
        return int(codigo[4:6]) in {1, 2, 3, 4}
    except ValueError:
        return False


def nombre_video_cmsv6(
    archivo: dict,
    device_id: str,
    fecha: datetime.date | datetime.datetime | None = None,
) -> str:
    file_path = str(archivo.get("file", "") or "")
    if file_path:
        nombre_original = Path(os.path.basename(file_path)).stem + ".mp4"
        if _nombre_compatible_mdvr(nombre_original):
            return nombre_original
    channel = int(archivo.get("chn", 0) or 0)
    begin = _cmsv6_segundos_a_hhmmss(archivo.get("beg", 0))
    end = _cmsv6_segundos_a_hhmmss(archivo.get("end", 0))
    fecha_nombre = _fecha_nombre_video(fecha)
    codigo = _codigo_video_compatible_mdvr(channel)
    return f"{device_id}-{fecha_nombre.strftime('%y%m%d')}-{begin:06d}-{end:06d}-{codigo}.mp4"


def _video_size_bytes(archivo: dict) -> int:
    try:
        return int(archivo.get("len", 0) or 0)
    except (TypeError, ValueError):
        return 0


def _video_channel_idx(archivo: dict) -> int:
    try:
        return int(archivo.get("chn", 0) or 0)
    except (TypeError, ValueError):
        return 0


def _segundos_video(archivo: dict, key: str) -> int:
    try:
        return max(0, int(archivo.get(key, 0) or 0))
    except (TypeError, ValueError):
        return 0


def _metadata_cmsv6_video(
    archivo: dict,
    *,
    fecha: datetime.date | datetime.datetime,
    nombre_mp4: str,
    descarga_usada: str | None = None,
) -> dict:
    beg = _segundos_video(archivo, "beg")
    end = _segundos_video(archivo, "end")
    metadata = {
        "fuente": "cmsv6",
        "archivo": nombre_mp4,
        "fecha_consulta": _fecha_nombre_video(fecha).isoformat(),
        "beg_segundo_dia": beg,
        "end_segundo_dia": end,
        "duracion_api_segundos": max(0, end - beg),
        "dev_idno": str(archivo.get("devIdno", archivo.get("DevIDNO", "")) or ""),
        "camara_indice": _video_channel_idx(archivo),
        "camara": _video_channel_idx(archivo) + 1,
        "loc": archivo.get("loc"),
        "svr": archivo.get("svr"),
        "file": str(archivo.get("file", "") or ""),
        "len": _video_size_bytes(archivo),
        "mediaType": archivo.get("mediaType"),
        "type": archivo.get("type"),
        "stream": archivo.get("stream"),
        "streamType": archivo.get("streamType"),
        "sourceId": archivo.get("sourceId"),
        "descarga_usada": descarga_usada or "",
    }
    return {key: value for key, value in metadata.items() if value not in (None, "")}


def _guardar_metadata_cmsv6_video(
    dest_mp4: Path,
    archivo: dict,
    *,
    fecha: datetime.date | datetime.datetime,
    nombre_mp4: str,
    descarga_usada: str | None = None,
) -> None:
    metadata = _metadata_cmsv6_video(
        archivo,
        fecha=fecha,
        nombre_mp4=nombre_mp4,
        descarga_usada=descarga_usada,
    )
    sidecar = Path(f"{dest_mp4}.cmsv6.json")
    tmp_sidecar = sidecar.with_suffix(sidecar.suffix + ".tmp")
    tmp_sidecar.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    tmp_sidecar.replace(sidecar)


def _completar_down_url_cmsv6(
    url: str,
    archivo: dict,
    *,
    fecha: datetime.date | datetime.datetime,
    nombre_mp4: str,
) -> str:
    if not url or _cmsv6_url_int_param(url, "DownType") != 3:
        return url
    try:
        archivo_loc = int(archivo.get("loc", 0) or 0)
    except (TypeError, ValueError):
        archivo_loc = 0
    url_fileloc = _cmsv6_url_int_param(url, "FILELOC")
    if url_fileloc not in (2, None) and archivo_loc != 2:
        return url

    fecha_nombre = _fecha_nombre_video(fecha)
    chn_mask = archivo.get("chnMask")
    if chn_mask in (None, ""):
        chn_mask = _video_channel_idx(archivo)
    fileattr = (
        archivo.get("FILEATTR")
        or archivo.get("fileattr")
        or archivo.get("fileAttr")
        or (2 if archivo_loc == 2 or url_fileloc == 2 else None)
    )
    params = {
        "SAVENAME": nombre_mp4,
        "YEAR": fecha_nombre.year % 100,
        "MON": fecha_nombre.month,
        "DAY": fecha_nombre.day,
        "BEG": _segundos_video(archivo, "beg"),
        "END": _segundos_video(archivo, "end"),
        "CHNMASK": chn_mask,
        "FILEATTR": fileattr,
    }
    return _cmsv6_url_completar_params(url, params, llenar_vacios=("SAVENAME",))


def _hhmmss_desde_segundos(segundos: int) -> str:
    segundos = max(0, int(segundos or 0)) % 86400
    hh = segundos // 3600
    mm = (segundos % 3600) // 60
    ss = segundos % 60
    return f"{hh:02d}:{mm:02d}:{ss:02d}"


def _tipo_turno_para_segundos(segundos: int) -> str:
    hora = (max(0, int(segundos or 0)) % 86400) // 3600
    if hora < 8:
        return "noche"
    if hora < 16:
        return "manana"
    return "tarde"


def _turno_label(tipo_turno: str) -> str:
    return {
        "noche": "Noche",
        "manana": "Dia",
        "tarde": "Tarde",
    }.get(tipo_turno or "", str(tipo_turno or "Turno"))


def _contexto_video_descarga(
    archivo: dict,
    *,
    config: CMSV6Config,
    fecha: datetime.datetime,
    index: int,
    total: int,
    nombre_mp4: str,
    dest_tmp: Path,
    dest_mp4: Path,
) -> dict:
    inicio = _segundos_video(archivo, "beg")
    fin = _segundos_video(archivo, "end")
    tipo_turno = _tipo_turno_para_segundos(inicio)
    total_bytes = _video_size_bytes(archivo)
    return {
        "fecha": fecha.date().isoformat(),
        "turno": tipo_turno,
        "turno_label": _turno_label(tipo_turno),
        "camara": _video_channel_idx(archivo) + 1,
        "video_index": index,
        "videos_total": total,
        "archivo": nombre_mp4,
        "equipo": config.device_id,
        "hora_inicio": _hhmmss_desde_segundos(inicio),
        "hora_fin": _hhmmss_desde_segundos(fin),
        "bytes_total": total_bytes,
        "mb_total": round(total_bytes / 1048576, 1) if total_bytes else None,
        "ruta_parcial": str(dest_tmp),
        "ruta_destino": str(dest_mp4),
    }


def _emit_progress(set_progress, progress, message, extra=None):
    if extra:
        try:
            set_progress(progress, message, extra)
            return
        except TypeError:
            pass
    set_progress(progress, message)


class _DayDownloadProgress:
    def __init__(self, set_progress, *, base_pct: int, end_pct: int, total: int):
        self.set_progress = set_progress
        self.total = max(1, int(total or 1))
        self.v_pct_base = base_pct + 18
        self.v_pct_span = max(1, end_pct - base_pct - 18)
        self.file_pcts: dict[int, int] = {}
        self.active: dict[int, dict] = {}
        self.last_global = self.v_pct_base
        self.lock = threading.Lock()

    def _global_pct_locked(self) -> int:
        unidades = sum(max(0, min(100, pct)) for pct in self.file_pcts.values()) / 100
        global_pct = self.v_pct_base + int((unidades / self.total) * self.v_pct_span)
        global_pct = min(max(global_pct, self.last_global), 99)
        self.last_global = global_pct
        return global_pct

    def update(
        self,
        index: int,
        contexto: dict,
        *,
        pct_file: int = 0,
        estado: str = "descargando",
        downloaded: int | None = None,
        total_bytes: int | None = None,
        error: str = "",
        active: bool = True,
    ) -> int:
        with self.lock:
            pct_file = int(max(0, min(100, pct_file)))
            self.file_pcts[index] = max(self.file_pcts.get(index, 0), pct_file)
            actual = dict(contexto)
            actual.update(
                {
                    "estado": estado,
                    "porcentaje_archivo": pct_file,
                    "bytes_descargados": downloaded,
                    "bytes_total": total_bytes or contexto.get("bytes_total"),
                    "mb_descargados": (
                        round(downloaded / 1048576, 1) if downloaded is not None else None
                    ),
                    "error": error,
                }
            )
            if active:
                self.active[index] = actual
            else:
                self.active.pop(index, None)

            global_pct = self._global_pct_locked()
            message = self._message(global_pct, actual)
            extra = {
                "descarga_actual": actual,
                "descargas_activas": [
                    self.active[key] for key in sorted(self.active)
                ],
            }
            _emit_progress(self.set_progress, global_pct, message, extra)
            return global_pct

    @staticmethod
    def _message(global_pct: int, contexto: dict) -> str:
        prefix = (
            f"[{global_pct}%] {contexto['fecha']} {contexto['turno_label']} "
            f"CH{contexto['camara']} video "
            f"{contexto['video_index']}/{contexto['videos_total']}: "
        )
        if contexto.get("estado") == "descargando":
            downloaded = contexto.get("mb_descargados")
            total = contexto.get("mb_total")
            if downloaded is not None and total:
                return f"{prefix}{contexto['porcentaje_archivo']}% | {downloaded:.1f}/{total:.1f} MB"
            return f"{prefix}{contexto['porcentaje_archivo']}%"
        return f"{prefix}{contexto.get('estado') or 'preparando'}"


def _filtrar_videos_testing(archivos, opts, require_downloadable=False):
    channel = str(opts.get("test_channel", "Todos"))
    max_mb = float(opts.get("test_max_mb", 0) or 0)
    result = list(archivos or [])
    if channel.startswith("CH"):
        try:
            channel_idx = int(channel[2:]) - 1
            result = [item for item in result if _video_channel_idx(item) == channel_idx]
        except ValueError:
            pass
    if max_mb > 0:
        max_bytes = max_mb * 1048576
        result = [item for item in result if _video_size_bytes(item) <= max_bytes]
    if require_downloadable:
        result = [
            item
            for item in result
            if str(item.get("file", "") or "").strip()
            or str(item.get("DownUrl", "") or "").strip()
            or str(item.get("PlaybackUrl", "") or "").strip()
        ]
    return result


def _ordenar_videos_testing(archivos, test_order):
    if test_order == "size_asc":
        return sorted(archivos, key=_video_size_bytes)
    if test_order == "size_desc":
        return sorted(archivos, key=_video_size_bytes, reverse=True)
    return list(archivos)


def _dias_en_rango(fecha_ini: datetime.datetime, fecha_fin: datetime.datetime):
    dias = []
    day = fecha_ini.date()
    while day <= fecha_fin.date():
        dias.append(day)
        day += datetime.timedelta(days=1)
    return dias


def _seleccionar_video_extremo(session, dias, mode, opts, config, log_fn, set_progress, label):
    best_arch = None
    best_day = None
    scan_days = list(reversed(dias)) if config.test_30d_scan_newest_first else list(dias)
    for index, day in enumerate(scan_days, start=1):
        pct = min(5 + int((index / max(1, len(scan_days))) * 13), 18)
        set_progress(pct, f"[{pct}%] {label}: revisando {day.isoformat()}...")
        log_fn(f"  [{label}] {index}/{len(scan_days)} {day.isoformat()}: consultando...")
        files = []
        for attempt in range(1, config.test_30d_day_retries + 2):
            try:
                files = session.get_video_files(datetime.datetime.combine(day, datetime.time.min), log_fn=log_fn)
            except Exception as exc:
                log_fn(f"    [{label}] consulta fallo: {exc}")
                files = []
            candidates = _filtrar_videos_testing(files, opts, require_downloadable=True)
            candidates = [
                item for item in candidates if _video_size_bytes(item) >= config.small_response_bytes
            ]
            if candidates or attempt > config.test_30d_day_retries:
                break
            if config.test_30d_day_retry_wait_secs:
                time.sleep(config.test_30d_day_retry_wait_secs)
        if not candidates:
            log_fn(f"    [{label}] 0 candidatos utiles")
            continue
        candidate = min(candidates, key=_video_size_bytes) if mode == "min" else max(candidates, key=_video_size_bytes)
        better = best_arch is None or (
            _video_size_bytes(candidate) < _video_size_bytes(best_arch)
            if mode == "min"
            else _video_size_bytes(candidate) > _video_size_bytes(best_arch)
        )
        log_fn(
            f"    [{label}] mejor del dia: CH{_video_channel_idx(candidate) + 1} | "
            f"{_video_size_bytes(candidate) / 1048576:.1f} MB | "
            f"{nombre_video_cmsv6(candidate, config.device_id, day)}"
        )
        if better:
            best_arch = dict(candidate)
            best_arch["_test_scan_date"] = day.isoformat()
            best_day = day
    if best_arch and best_day:
        log_fn(
            f"[{label}] Seleccion final: {best_day.isoformat()} | "
            f"CH{_video_channel_idx(best_arch) + 1} | "
            f"{_video_size_bytes(best_arch) / 1048576:.1f} MB | "
            f"{nombre_video_cmsv6(best_arch, config.device_id, best_day)}"
        )
        return {best_day: [best_arch]}, [best_day]
    log_fn(f"[{label}] No se encontro un video descargable con esos filtros.")
    return {}, []


def _descarga_suficiente(path: Path, file_len: int, config: CMSV6Config) -> bool:
    size = path.stat().st_size if path.exists() else 0
    if file_len > config.small_response_bytes:
        min_size = int(file_len * 0.95)
        max_size = max(file_len + config.small_response_bytes, int(file_len * 1.05))
        return min_size <= size <= max_size
    return size >= config.small_response_bytes


def _descargar_video_archivo(
    session,
    archivo,
    *,
    fecha,
    carpeta_dia,
    log_fn,
    tracker: _DayDownloadProgress,
    config,
    index: int,
    total: int,
) -> dict:
    resumen = {"descargados": 0, "omitidos": 0, "errores": 0}
    nombre_mp4 = nombre_video_cmsv6(archivo, config.device_id, fecha)
    dest_mp4 = carpeta_dia / nombre_mp4
    dest_tmp = carpeta_dia / (Path(nombre_mp4).stem + ".tmp")
    contexto = _contexto_video_descarga(
        archivo,
        config=config,
        fecha=fecha,
        index=index,
        total=total,
        nombre_mp4=nombre_mp4,
        dest_tmp=dest_tmp,
        dest_mp4=dest_mp4,
    )
    channel = contexto["camara"]
    size_mb = _video_size_bytes(archivo) / 1048576
    expected_secs = None
    try:
        expected_secs = int(archivo.get("end", 0) or 0) - int(archivo.get("beg", 0) or 0)
        if expected_secs <= 0:
            expected_secs = None
    except Exception:
        expected_secs = None

    if dest_mp4.exists() and dest_mp4.stat().st_size > 4096 and is_mp4(dest_mp4):
        _guardar_metadata_cmsv6_video(
            dest_mp4,
            archivo,
            fecha=fecha,
            nombre_mp4=nombre_mp4,
            descarga_usada="existente",
        )
        log_fn(
            f"  [{index}/{total}] EXISTE {contexto['fecha']} "
            f"{contexto['turno_label']} CH{channel}: {nombre_mp4}"
        )
        tracker.update(index, contexto, pct_file=100, estado="ya existe", active=False)
        resumen["omitidos"] += 1
        return resumen

    fpath = str(archivo.get("file", "") or "").strip()
    down_task_url = str(archivo.get("DownTaskUrl", "") or "").strip()
    raw_down_url = str(archivo.get("DownUrl", "") or "").strip()
    raw_down_url = _completar_down_url_cmsv6(
        raw_down_url,
        archivo,
        fecha=fecha,
        nombre_mp4=nombre_mp4,
    )
    raw_play_url = str(archivo.get("PlaybackUrl", "") or "").strip()
    file_len = int(archivo.get("len", 0) or 0)
    if not (fpath or raw_down_url or raw_play_url):
        log_fn(f"  [{index}/{total}] SIN datos de descarga: {nombre_mp4}")
        tracker.update(
            index,
            contexto,
            pct_file=100,
            estado="error",
            error="sin datos de descarga",
            active=False,
        )
        resumen["errores"] += 1
        return resumen

    log_fn(
        f"\n  [{index}/{total}] {contexto['fecha']} {contexto['turno_label']} "
        f"CH{channel} | {contexto['hora_inicio']}-{contexto['hora_fin']} | "
        f"{nombre_mp4} | {size_mb:.1f} MB"
    )
    tracker.update(index, contexto, pct_file=0, estado="preparando")

    if down_task_url:
        try:
            task_url = session.refresh_url(down_task_url).replace(" ", "%20")
            request = urllib.request.Request(task_url, method="GET")
            request.add_header("User-Agent", "Mozilla/5.0 (Windows NT 10.0; Win64; x64)")
            with session.opener.open(request, timeout=20) as response:
                task_response = json.loads(response.read())
            log_fn(
                f"    Tarea CMSV6: result={task_response.get('result', '?')} "
                f"id={task_response.get('taskId', task_response.get('id', '?'))}"
            )
            if _result_code(task_response) in (None, 0):
                task_info = session.wait_download_task(archivo, fecha, log_fn)
                if task_info:
                    status = session._download_task_status(task_info)
                    if status == 4:
                        log_fn("    Tarea CMSV6 lista para descarga.")
                    else:
                        log_fn(f"    Tarea CMSV6 no quedo lista; status={status}.")
        except Exception as exc:
            log_fn(f"    Tarea CMSV6 ignorada por error: {exc}")

    dispositivo_online = None
    try:
        gps = session.get_gps()
        if gps:
            dispositivo_online = gps[0].get("ol") in (1, "1", True)
    except Exception:
        pass

    candidates = []
    if dispositivo_online is False:
        if raw_play_url:
            candidates.append(("PlaybackUrl", raw_play_url, config.playback_stall_secs))
        if raw_down_url:
            candidates.append(("DownUrl", raw_down_url, config.downurl_stall_secs))
        log_fn("    Dispositivo OFFLINE: usando PlaybackUrl primero")
    else:
        if raw_down_url:
            candidates.append(("DownUrl", raw_down_url, config.downurl_stall_secs))
        if raw_play_url:
            candidates.append(("PlaybackUrl", raw_play_url, config.playback_stall_secs))
        if dispositivo_online:
            log_fn("    Dispositivo ONLINE: usando DownUrl primero")

    download_ok = False
    last_error = "sin intento"
    download_label = ""

    def progress(downloaded, total_bytes):
        pct_file = downloaded * 100 // total_bytes if total_bytes else 0
        tracker.update(
            index,
            contexto,
            pct_file=pct_file,
            estado="descargando",
            downloaded=downloaded,
            total_bytes=total_bytes,
        )

    for round_idx in range(1, config.url_rounds + 1):
        if download_ok:
            break
        if round_idx > 1:
            try:
                session.full_relogin()
            except Exception as exc:
                log_fn(f"    Re-login para ronda {round_idx} fallo: {exc}")
        for label, raw_url, stall_secs in candidates:
            try:
                if dest_tmp.exists() and dest_tmp.stat().st_size < config.small_response_bytes:
                    dest_tmp.unlink()
                log_fn(f"    [{label}] descargando...")
                session.download_file(
                    raw_url,
                    dest_tmp,
                    progress_cb=progress,
                    log_fn=log_fn,
                    max_retries=3,
                    dl_timeout=max(120, min(int(file_len / (100 * 1024)) if file_len else 120, 1800)),
                    stall_secs=stall_secs,
                )
                if dest_tmp.exists() and _descarga_suficiente(dest_tmp, file_len, config):
                    download_ok = True
                    download_label = label
                    break
                last_error = "descarga incompleta o vacia"
                log_fn(f"    [{label}] {last_error}; probando otra URL...")
            except (StalledDownloadError, SlowDownloadError, DownloadTimeLimitError) as exc:
                last_error = str(exc)
                log_fn(f"    [{label}] {exc}; probando otra URL...")
            except Exception as exc:
                last_error = str(exc)
                log_fn(f"    Error [{label}]: {exc}")
        if not download_ok and round_idx < config.url_rounds and config.url_round_wait_secs:
            log_fn(f"    Sin exito en ronda {round_idx}; esperando {config.url_round_wait_secs}s...")
            time.sleep(config.url_round_wait_secs)

    try:
        if not download_ok:
            raise Exception(f"Ninguna URL funciono. Ultimo error: {last_error}")

        raw4 = b""
        if dest_tmp.exists() and dest_tmp.stat().st_size >= 4:
            with open(dest_tmp, "rb") as archivo_tmp:
                raw4 = archivo_tmp.read(8)
        fmt_hex = raw4[:4].hex() if raw4 else "????"
        log_fn(f"    Header: {fmt_hex} | Tamaño: {dest_tmp.stat().st_size / 1048576:.2f} MB")

        if is_mp4(dest_tmp):
            if dest_mp4.exists():
                dest_mp4.unlink()
            dest_tmp.rename(dest_mp4)
            _guardar_metadata_cmsv6_video(
                dest_mp4,
                archivo,
                fecha=fecha,
                nombre_mp4=nombre_mp4,
                descarga_usada=download_label,
            )
            log_fn(f"    MP4 nativo directo: {nombre_mp4}")
            tracker.update(index, contexto, pct_file=100, estado="completo", active=False)
            resumen["descargados"] += 1
            return resumen

        tracker.update(index, contexto, pct_file=99, estado="convirtiendo")
        fuente_original = str(archivo.get("file") or "").lower()
        raw_h264 = (
            fuente_original.endswith(".h264")
            or fuente_original.endswith(".264")
            or is_h264(dest_tmp)
        )
        if convert_to_mp4(
            dest_tmp,
            dest_mp4,
            log_fn,
            expected_secs=expected_secs,
            raw_h264=raw_h264,
        ):
            dest_tmp.unlink(missing_ok=True)
            _guardar_metadata_cmsv6_video(
                dest_mp4,
                archivo,
                fecha=fecha,
                nombre_mp4=nombre_mp4,
                descarga_usada=f"{download_label}:convertido" if download_label else "convertido",
            )
            log_fn(f"    MP4 OK: {nombre_mp4} ({dest_mp4.stat().st_size / 1048576:.1f} MB)")
            tracker.update(index, contexto, pct_file=100, estado="completo", active=False)
            resumen["descargados"] += 1
        else:
            raw_dest = carpeta_dia / (Path(nombre_mp4).stem + f"_{fmt_hex}.raw")
            if dest_tmp.exists():
                dest_tmp.rename(raw_dest)
            log_fn(f"    Conversion fallo; crudo guardado: {raw_dest.name}")
            tracker.update(
                index,
                contexto,
                pct_file=100,
                estado="error",
                error="conversion fallo",
                active=False,
            )
            resumen["errores"] += 1
    except Exception as exc:
        log_fn(f"    ERROR: {exc}")
        tracker.update(
            index,
            contexto,
            pct_file=100,
            estado="error",
            error=str(exc),
            active=False,
        )
        resumen["errores"] += 1
        if dest_mp4.exists() and not is_mp4(dest_mp4):
            dest_mp4.unlink()
    return resumen


def _merge_resumen_descarga(destino: dict, parcial: dict):
    destino["descargados"] += int(parcial.get("descargados", 0) or 0)
    destino["omitidos"] += int(parcial.get("omitidos", 0) or 0)
    destino["errores"] += int(parcial.get("errores", 0) or 0)


def _descargar_videos_dia(
    session,
    archivos,
    fecha,
    carpeta_videos_base,
    log_fn,
    set_progress,
    config,
    base_pct,
    end_pct,
    download_workers=1,
    on_file_complete=None,
):
    resumen = {"descargados": 0, "omitidos": 0, "errores": 0, "total": len(archivos or [])}
    if not archivos:
        log_fn("  Sin videos para este dia.")
        return resumen

    carpeta_dia = carpeta_videos_base / fecha.strftime("%Y-%m-%d")
    carpeta_dia.mkdir(parents=True, exist_ok=True)
    total = len(archivos)
    tracker = _DayDownloadProgress(set_progress, base_pct=base_pct, end_pct=end_pct, total=total)
    workers = max(1, min(int(download_workers or 1), total))
    log_fn(f"  Encontrados: {total} archivos en CMSV6")
    if workers > 1 and getattr(config, "serial_downloads_per_device", True):
        log_fn(
            "  Descargas serializadas por equipo MDVR: "
            f"se usara 1 video a la vez (configurado: {workers})"
        )
        workers = 1
    if workers > 1:
        log_fn(f"  Descargas paralelas: {workers} videos a la vez")

    if workers == 1:
        for index, archivo in enumerate(archivos, start=1):
            parcial = _descargar_video_archivo(
                session,
                archivo,
                fecha=fecha,
                carpeta_dia=carpeta_dia,
                log_fn=log_fn,
                tracker=tracker,
                config=config,
                index=index,
                total=total,
            )
            _merge_resumen_descarga(resumen, parcial)
            if on_file_complete:
                try:
                    on_file_complete(fecha, archivo, parcial)
                except Exception as exc:
                    log_fn(f"    ERROR notificando archivo completado: {exc}")
    else:
        def _run(index, archivo):
            worker_session = CMSV6Session(config)
            worker_session.login()
            parcial = _descargar_video_archivo(
                worker_session,
                archivo,
                fecha=fecha,
                carpeta_dia=carpeta_dia,
                log_fn=log_fn,
                tracker=tracker,
                config=config,
                index=index,
                total=total,
            )
            return archivo, parcial

        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = [
                executor.submit(_run, index, archivo)
                for index, archivo in enumerate(archivos, start=1)
            ]
            for future in as_completed(futures):
                try:
                    archivo, parcial = future.result()
                    _merge_resumen_descarga(resumen, parcial)
                    if on_file_complete:
                        try:
                            on_file_complete(fecha, archivo, parcial)
                        except Exception as exc:
                            log_fn(f"    ERROR notificando archivo completado: {exc}")
                except Exception as exc:
                    log_fn(f"    ERROR descarga paralela: {exc}")
                    resumen["errores"] += 1

    log_fn(
        f"\n  VIDEOS COMPLETADOS: {resumen['descargados']} descargados | "
        f"{resumen['omitidos']} ya existian | {resumen['errores']} errores"
    )
    return resumen


def ejecutar_rango(
    carpeta_base,
    fecha_ini,
    fecha_fin,
    log_fn,
    set_progress,
    opts=None,
    config=None,
    on_day_complete=None,
    on_file_complete=None,
):
    config = config or CMSV6Config.from_settings(carpeta_base)
    config.validate()
    opts = opts or {}
    carpeta_base = Path(carpeta_base or config.output_dir)
    carpeta_base.mkdir(parents=True, exist_ok=True)
    carpeta_videos_base = carpeta_base / f"{config.device_id}({config.device_id})"

    hacer_ruta = bool(opts.get("excel_ruta", True))
    hacer_alarmas = bool(opts.get("excel_alarmas", True))
    hacer_excel = hacer_ruta or hacer_alarmas
    hacer_videos = bool(opts.get("videos", True))
    test_30d_mode = str(opts.get("test_30d_mode", "off") or "off")
    test_range_mode = str(opts.get("test_range_mode", "off") or "off")
    try:
        download_workers = max(
            1,
            int(opts.get("download_workers") or getattr(config, "download_workers", 1)),
        )
    except (TypeError, ValueError):
        download_workers = 1
    if test_30d_mode not in ("off", "min", "max"):
        test_30d_mode = "off"
    if test_range_mode not in ("off", "min", "max"):
        test_range_mode = "off"
    if test_30d_mode in ("min", "max"):
        test_range_mode = "off"

    dias = _dias_en_rango(fecha_ini, fecha_fin)
    if opts.get("testing") and test_30d_mode in ("min", "max"):
        end_day = fecha_fin.date()
        start_day = end_day - datetime.timedelta(days=29)
        dias = [start_day + datetime.timedelta(days=index) for index in range(30)]
        fecha_ini = datetime.datetime.combine(start_day, datetime.time.min)
        fecha_fin = datetime.datetime.combine(end_day, datetime.time.max.replace(microsecond=0))

    log_fn("\n" + "=" * 56)
    log_fn(f"  {fecha_ini.strftime('%Y-%m-%d %H:%M')} -> {fecha_fin.strftime('%Y-%m-%d %H:%M')}")
    log_fn(f"  Dias: {len(dias)}")
    log_fn(f"  Salida: {carpeta_base}")
    log_fn("=" * 56)

    set_progress(2, "[2%] Conectando al servidor CMSV6...")
    session = CMSV6Session(config)
    session.login()
    log_fn("Login CMSV6 exitoso.")
    set_progress(5, "[5%] Login OK")

    test_target_by_day = None
    if hacer_videos and opts.get("testing") and test_30d_mode in ("min", "max"):
        test_target_by_day, dias = _seleccionar_video_extremo(
            session, dias, test_30d_mode, opts, config, log_fn, set_progress, "TEST 30D"
        )
    elif hacer_videos and opts.get("testing") and test_range_mode in ("min", "max"):
        test_target_by_day, dias = _seleccionar_video_extremo(
            session, dias, test_range_mode, opts, config, log_fn, set_progress, "TEST RANGO"
        )

    resumen_total = {
        "dias": len(dias),
        "videos_descargados": 0,
        "videos_omitidos": 0,
        "videos_errores": 0,
        "salida": str(carpeta_base),
    }

    for day_index, day in enumerate(dias, start=1):
        fecha = datetime.datetime.combine(day, datetime.time.min)
        base_pct = 5 + int(((day_index - 1) / max(1, len(dias))) * 90)
        end_pct = 5 + int((day_index / max(1, len(dias))) * 90)
        day_label = fecha.strftime("%Y-%m-%d")
        log_fn("\n" + "=" * 56)
        log_fn(f"  Dia {day_index}/{len(dias)}: {day_label}")
        log_fn("=" * 56)

        gps_data = {}
        if day == datetime.datetime.now().date() or day_index == len(dias):
            set_progress(base_pct + 2, f"[{base_pct + 2}%] {day_label}: GPS...")
            try:
                gps_list = session.get_gps()
                if gps_list:
                    gps = gps_list[0]
                    gps_data = {
                        "Fecha/Hora": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "ID Dispositivo": gps.get("id", ""),
                        "Latitud": float(gps.get("mlat", 0)),
                        "Longitud": float(gps.get("mlng", 0)),
                        "Velocidad km/h": _normalizar_vel(gps.get("sp", 0)),
                        "Direccion °": gps.get("hx", 0),
                        "En linea": "Si" if gps.get("ol") else "No",
                        "Hora GPS": gps.get("gt", ""),
                        "Hora servidor": gps.get("rt", ""),
                        "Señal red": gps.get("net", ""),
                        "Kilometraje": gps.get("mileage", ""),
                        "Estado": gps.get("status", ""),
                        "raw": gps,
                    }
                    log_fn(
                        f"  GPS OK: Lat={gps_data['Latitud']:.5f} | "
                        f"Lng={gps_data['Longitud']:.5f} | Online={gps_data['En linea']}"
                    )
                else:
                    log_fn("  Sin datos GPS.")
            except Exception as exc:
                log_fn(f"  ERROR GPS: {exc}")

        set_progress(base_pct + 7, f"[{base_pct + 7}%] {day_label}: ruta GPS...")
        try:
            tracks = session.get_track(fecha)
            log_fn(f"  Ruta: {len(tracks)} puntos obtenidos")
            guardar_tracks_json(carpeta_base, config.device_id, day, tracks, log_fn)
        except Exception as exc:
            tracks = []
            log_fn(f"  ERROR ruta: {exc}")

        set_progress(base_pct + 12, f"[{base_pct + 12}%] {day_label}: alarmas...")
        try:
            alarmas = session.get_alarms(fecha)
            log_fn(f"  Alarmas: {len(alarmas)} obtenidas")
        except Exception as exc:
            alarmas = []
            log_fn(f"  ERROR alarmas: {exc}")

        eventos = []
        if hacer_excel and tracks:
            set_progress(base_pct + 15, f"[{base_pct + 15}%] {day_label}: eventos...")
            try:
                eventos = detectar_eventos_ruta(tracks, config.device_id)
                log_fn(f"  Eventos detectados: {len(eventos)}")
            except Exception as exc:
                log_fn(f"  ERROR eventos: {exc}")

        if hacer_excel and (gps_data or tracks or alarmas or eventos):
            set_progress(base_pct + 18, f"[{base_pct + 18}%] {day_label}: generando Excel...")
            try:
                exportar_excel(
                    carpeta_base,
                    config.device_id,
                    datetime.datetime.combine(day, datetime.time.min),
                    datetime.datetime.combine(day, datetime.time.max.replace(microsecond=0)),
                    gps_data,
                    tracks,
                    alarmas,
                    eventos,
                    log_fn,
                    hacer_ruta=hacer_ruta,
                    hacer_alarmas=hacer_alarmas,
                )
            except Exception as exc:
                log_fn(f"  ERROR Excel: {exc}")
        elif not hacer_excel:
            log_fn("  Excel omitido por opciones.")

        if not hacer_videos:
            log_fn("  Videos omitidos por opciones.")
            if on_day_complete:
                try:
                    on_day_complete(
                        day,
                        {
                            "descargados": 0,
                            "omitidos": 0,
                            "errores": 0,
                            "total": 0,
                        },
                    )
                except Exception as exc:
                    log_fn(f"  ERROR notificando cierre de dia {day_label}: {exc}")
            continue

        set_progress(base_pct + 18, f"[{base_pct + 18}%] {day_label}: buscando videos...")
        if test_target_by_day is not None:
            archivos = test_target_by_day.get(day, [])
        else:
            archivos = session.get_video_files(fecha, log_fn=log_fn)
            if archivos and opts.get("testing"):
                total_servidor = len(archivos)
                archivos = _filtrar_videos_testing(archivos, opts)
                archivos = _ordenar_videos_testing(archivos, str(opts.get("test_order", "size_asc")))
                archivos = archivos[: max(1, int(opts.get("test_limit", 1) or 1))]
                log_fn(f"  [TEST] Seleccionados: {len(archivos)}/{total_servidor} clips")

        resumen_dia = _descargar_videos_dia(
            session,
            archivos,
            fecha,
            carpeta_videos_base,
            log_fn,
            set_progress,
            config,
            base_pct,
            end_pct,
            download_workers=download_workers,
            on_file_complete=on_file_complete,
        )
        resumen_total["videos_descargados"] += resumen_dia["descargados"]
        resumen_total["videos_omitidos"] += resumen_dia["omitidos"]
        resumen_total["videos_errores"] += resumen_dia["errores"]
        if on_day_complete:
            try:
                on_day_complete(day, resumen_dia)
            except Exception as exc:
                log_fn(f"  ERROR notificando cierre de dia {day_label}: {exc}")

    set_progress(100, "Completado")
    log_fn(f"\n[OK] Todo en: {carpeta_base}\n")
    return resumen_total


def _parse_datetime(value: str, *, end_of_day=False):
    value = str(value or "").strip()
    if not value:
        raise ValidationError("Fecha/hora requerida.")
    if len(value) == 10:
        suffix = "23:59" if end_of_day else "00:00"
        value = f"{value} {suffix}"
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise ValidationError("Formato de fecha/hora inválido. Use YYYY-MM-DD HH:MM.")


def ejecutar_job_cmsv6(params, log_fn, set_progress):
    params = dict(params or {})
    config = CMSV6Config.from_settings(params.get("output_dir"))
    config.validate()
    output_dir = params.get("output_dir") or config.output_dir
    mode = str(params.get("modo", "dia") or "dia")
    opts = {
        "excel_ruta": bool(params.get("excel_ruta", True)),
        "excel_alarmas": bool(params.get("excel_alarmas", True)),
        "videos": bool(params.get("videos", True)),
        "testing": bool(params.get("testing", False)),
        "test_order": params.get("test_order", "size_asc"),
        "test_limit": int(params.get("test_limit", 1) or 1),
        "test_max_mb": float(params.get("test_max_mb", 0) or 0),
        "test_channel": params.get("test_channel", "Todos"),
        "test_30d_mode": params.get("test_30d_mode", "off"),
        "test_range_mode": params.get("test_range_mode", "off"),
    }

    results = []
    if mode == "dia":
        day = _parse_datetime(str(params.get("dia") or datetime.datetime.now().strftime("%Y-%m-%d")))
        fecha_ini = day.replace(hour=0, minute=0, second=0, microsecond=0)
        fecha_fin = day.replace(hour=23, minute=59, second=59, microsecond=0)
        results.append(ejecutar_rango(output_dir, fecha_ini, fecha_fin, log_fn, set_progress, opts, config))
    elif mode == "rango":
        fecha_ini = _parse_datetime(params.get("fecha_inicio"))
        fecha_fin = _parse_datetime(params.get("fecha_fin"), end_of_day=True)
        if fecha_ini > fecha_fin:
            fecha_ini, fecha_fin = fecha_fin, fecha_ini
        results.append(ejecutar_rango(output_dir, fecha_ini, fecha_fin, log_fn, set_progress, opts, config))
    elif mode == "auto":
        start_at = _parse_datetime(params.get("fecha_inicio") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
        interval_minutes = max(1, int(params.get("intervalo_minutos", 15) or 15))
        max_cycles = max(0, int(params.get("auto_ciclos", 1) or 0))
        cycle = 0
        while max_cycles == 0 or cycle < max_cycles:
            cycle += 1
            now = datetime.datetime.now()
            if cycle == 1 and start_at > now:
                wait_seconds = int((start_at - now).total_seconds())
                log_fn(f"Esperando inicio automatico: {start_at.strftime('%Y-%m-%d %H:%M')}")
                for _ in range(wait_seconds):
                    time.sleep(1)
            now = datetime.datetime.now()
            log_fn(f"\n[AUTO] Ciclo {cycle}: {now.strftime('%Y-%m-%d %H:%M')}")
            results.append(ejecutar_rango(output_dir, now, now, log_fn, set_progress, opts, config))
            if max_cycles and cycle >= max_cycles:
                break
            log_fn(f"[AUTO] Proxima ejecucion en {interval_minutes} minutos.")
            for _ in range(interval_minutes * 60):
                time.sleep(1)
    else:
        raise ValidationError("Modo CMSV6 inválido.")

    import_result = None
    if params.get("importar_django", True):
        from dashboard.services.importar_videos_mdvr import importar_videos_mdvr

        log_fn("Importando archivos descargados a modelos Django...")
        import_result = importar_videos_mdvr(base_dir=output_dir, importar_velocidades=True)
        log_fn(
            "Importacion Django finalizada: "
            f"{import_result.get('videos_creados', 0)} videos procesados."
        )

    return {"resultados_descarga": results, "importacion_django": import_result, "salida": output_dir}
